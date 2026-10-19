from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


class DatasetEvaluator:
    """Koordiniert alle Records nacheinander.

    Pro Record:
        laden -> Pipeline ausfuehren -> evaluieren -> naechster Record

    PipelineEvaluator bewertet einen einzigen Record.
    DatasetEvaluator findet und verarbeitet alle Records.
    """

    def __init__(
        self,
        data_loader,
        pipeline_evaluator,
        raw_data_path,
        output_file,
        channel=0,
    ):
        self.data_loader = data_loader
        self.pipeline_evaluator = pipeline_evaluator
        self.raw_data_path = Path(raw_data_path)
        self.output_file = Path(output_file)
        self.channel = channel

    def run(self):
        all_record_results = []
        pipeline_summaries = []

        folders = sorted(
            folder
            for folder in self.raw_data_path.iterdir()
            if folder.is_dir()
        )

        for folder in folders:
            folder_name = folder.name
            reference_name = (
                folder_name[:-1]
                if folder_name.endswith("e")
                else folder_name
            )

            record_names = sorted(
                file.stem for file in folder.glob("*.hea")
            )

            if reference_name not in record_names:
                print(
                    f"Uebersprungen: Referenz {reference_name} "
                    f"fehlt in {folder}"
                )
                continue

            for noisy_name in record_names:
                if noisy_name == reference_name:
                    continue

                print(f"Pipeline und Evaluation: {noisy_name}")

                reference_record = self.data_loader.load_record(
                    Path(folder_name) / reference_name
                )
                noisy_record = self.data_loader.load_record(
                    Path(folder_name) / noisy_name
                )

                clean_signal = self.data_loader.get_signal(
                    reference_record,
                    channel=self.channel,
                )
                noisy_signal = self.data_loader.get_signal(
                    noisy_record,
                    channel=self.channel,
                )
                sampling_rate = self.data_loader.get_sampling_rate(
                    reference_record
                )

                record_df, pipeline_result = (
                    self.pipeline_evaluator.evaluate(
                        clean_signal=clean_signal,
                        noisy_signal=noisy_signal,
                        sampling_rate=sampling_rate,
                        record_name=noisy_name,
                    )
                )

                all_record_results.append(record_df)
                pipeline_summaries.append(
                    {
                        "record": noisy_name,
                        "total_segments": pipeline_result.total_segments,
                        "kept_segments": pipeline_result.kept_segments,
                        "filtered_segments": pipeline_result.filtered_segments,
                        "rejected_segments": pipeline_result.rejected_segments,
                        "keep_rate": pipeline_result.keep_rate,
                        "filter_rate": pipeline_result.filter_rate,
                        "reject_rate": pipeline_result.reject_rate,
                        "coverage": pipeline_result.coverage,
                    }
                )

        if not all_record_results:
            print("Keine Records ausgewertet.")
            return None

        record_results_df = pd.concat(
            all_record_results,
            ignore_index=True,
        )
        pipeline_summary_df = pd.DataFrame(pipeline_summaries)
        overall_summary_df = self._build_overall_summary(
            record_results_df
        )

        self._save(
            overall_summary_df,
            record_results_df,
            pipeline_summary_df,
        )

        print(f"Gespeichert: {self.output_file}")
        return self.output_file

    @staticmethod
    def _build_overall_summary(record_results_df):
        numeric_columns = [
            "coverage",
            "keep_rate",
            "filter_rate",
            "reject_rate",
            "medical_score",
            "morphology_score",
            "fidelity_score",
            "overall_score",
            "worst_improvement",
            "negative_metric_count",
            "strong_negative_metric_count",
            "total_runtime_ms",
            "mean_runtime_ms_per_segment",
            "peak_memory_kb",
            "runtime_percent_of_always",
            "runtime_saving_percent",
            "memory_percent_of_always",
            "memory_saving_percent",
        ]

        aggregations = {
            column: "mean"
            for column in numeric_columns
            if column in record_results_df.columns
        }

        return (
            record_results_df.groupby("strategy", as_index=False)
            .agg(aggregations)
        )

    def _save(self, overall_df, records_df, pipeline_df):
        self.output_file.parent.mkdir(parents=True, exist_ok=True)

        with pd.ExcelWriter(
            self.output_file,
            engine="openpyxl",
        ) as writer:
            overall_df.to_excel(
                writer,
                sheet_name="overall_comparison",
                index=False,
            )
            records_df.to_excel(
                writer,
                sheet_name="record_comparison",
                index=False,
            )
            pipeline_df.to_excel(
                writer,
                sheet_name="pipeline_summary",
                index=False,
            )
            pd.DataFrame().to_excel(
                writer,
                sheet_name="charts",
                index=False,
            )

        workbook = load_workbook(self.output_file)
        self._format_workbook(workbook)
        self._add_charts(workbook)
        workbook.save(self.output_file)

    @staticmethod
    def _format_workbook(workbook):
        fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid",
        )
        font = Font(color="FFFFFF", bold=True)

        for index, worksheet in enumerate(workbook.worksheets, start=1):
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False

            for cell in worksheet[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(horizontal="center")

            for column_cells in worksheet.columns:
                letter = get_column_letter(column_cells[0].column)
                width = max(
                    [len(str(cell.value)) for cell in column_cells if cell.value is not None]
                    or [10]
                )
                worksheet.column_dimensions[letter].width = min(width + 2, 35)

            if worksheet.title == "charts" or worksheet.max_row < 2:
                continue

            table = Table(
                displayName=f"Table_{index}_{worksheet.title}"[:30],
                ref=(
                    f"A1:{get_column_letter(worksheet.max_column)}"
                    f"{worksheet.max_row}"
                ),
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9",
                showRowStripes=True,
                showFirstColumn=False,
                showLastColumn=False,
                showColumnStripes=False,
            )
            worksheet.add_table(table)

    @staticmethod
    def _add_charts(workbook):
        source = workbook["overall_comparison"]
        target = workbook["charts"]
        headers = {
            str(cell.value): cell.column
            for cell in source[1]
            if cell.value is not None
        }

        if "strategy" not in headers:
            return

        categories = Reference(
            source,
            min_col=headers["strategy"],
            min_row=2,
            max_row=source.max_row,
        )

        score_headers = [
            header
            for header in [
                "medical_score",
                "morphology_score",
                "fidelity_score",
                "overall_score",
            ]
            if header in headers
        ]

        if score_headers:
            chart = BarChart()
            chart.title = "RAW vs Bandpass vs Selective vs Oracle"
            chart.y_axis.title = "Mittlere Verbesserung"
            chart.x_axis.title = "Strategie"
            chart.style = 10
            chart.height = 12
            chart.width = 22

            for header in score_headers:
                data = Reference(
                    source,
                    min_col=headers[header],
                    min_row=1,
                    max_row=source.max_row,
                )
                chart.add_data(data, titles_from_data=True)

            chart.set_categories(categories)
            target.add_chart(chart, "A2")

        if "coverage" in headers:
            coverage_chart = BarChart()
            coverage_chart.title = "Coverage der Strategien"
            coverage_chart.y_axis.title = "Coverage"
            coverage_chart.x_axis.title = "Strategie"
            coverage_chart.style = 11
            coverage_chart.height = 12
            coverage_chart.width = 22

            data = Reference(
                source,
                min_col=headers["coverage"],
                min_row=1,
                max_row=source.max_row,
            )
            coverage_chart.add_data(data, titles_from_data=True)
            coverage_chart.set_categories(categories)
            target.add_chart(coverage_chart, "A26")
