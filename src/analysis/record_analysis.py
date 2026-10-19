from pathlib import Path

import numpy as np
import pandas as pd

from openpyxl.chart import Reference, ScatterChart, Series

from src.analysis.base_analysis import BaseAnalysis


class RecordAnalysis(BaseAnalysis):
    """
    Fuehrt fuer jede results-Datei eine separate Analyse durch.

    Ausgabe:
        output/<record>/analysis/
        results_<record>_improvements.xlsx
    """

    QUALITY_WEIGHTS = {
        "quality_estimated_snr_score": 0.40,
        "quality_rr_plausibility_score": 0.25,
        "quality_beat_template_correlation_score": 0.20,
        "quality_baseline_drift_score": 0.10,
        "quality_bsqi_score": 0.05,
    }

    N_SUMMARY_BINS = 5
    N_DECISION_BINS = 10

    def add_segment_quality_score(
        self,
        result_df: pd.DataFrame,
    ) -> pd.DataFrame:
        result_df = result_df.copy()
        score_columns = []

        for metric in self.QUALITY_METRICS:
            if metric not in result_df.columns:
                continue

            score_column = f"{metric}_score"

            percentile_rank = result_df[metric].rank(
                pct=True,
                method="average",
                na_option="keep",
            )

            if metric in self.HIGHER_IS_BETTER_QUALITY_METRICS:
                result_df[score_column] = 1.0 - percentile_rank
            elif metric in self.LOWER_IS_BETTER_QUALITY_METRICS:
                result_df[score_column] = percentile_rank
            else:
                raise ValueError(
                    "Keine Qualitaetsrichtung definiert fuer: "
                    f"{metric}"
                )

            score_columns.append(score_column)

        if len(score_columns) == 0:
            insert_position = min(3, len(result_df.columns))
            result_df.insert(
                insert_position,
                "segment_quality_score",
                np.nan,
            )
            return result_df

        weighted_sum = pd.Series(
            0.0,
            index=result_df.index,
            dtype=float,
        )

        available_weight = pd.Series(
            0.0,
            index=result_df.index,
            dtype=float,
        )

        for score_column in score_columns:
            weight = self.QUALITY_WEIGHTS.get(score_column, 0.0)

            if weight <= 0.0:
                continue

            valid_mask = result_df[score_column].notna()

            weighted_values = np.multiply(
                result_df.loc[valid_mask, score_column],
                weight,
            )

            weighted_sum.loc[valid_mask] += weighted_values
            available_weight.loc[valid_mask] += weight

        segment_quality_score = pd.Series(
            np.nan,
            index=result_df.index,
            dtype=float,
        )

        valid_weight_mask = available_weight > self.EPSILON

        segment_quality_score.loc[valid_weight_mask] = (
            weighted_sum.loc[valid_weight_mask]
            / available_weight.loc[valid_weight_mask]
        )

        insert_position = min(3, len(result_df.columns))

        result_df.insert(
            insert_position,
            "segment_quality_score",
            segment_quality_score,
        )

        return result_df

    def build_improvement_table(
        self,
        df: pd.DataFrame,
        input_file: Path,
    ) -> pd.DataFrame:
        results = []

        if "segment_id" not in df.columns:
            print(f"FEHLER: segment_id fehlt in {input_file}")
            return pd.DataFrame()

        if "method" not in df.columns:
            print(f"FEHLER: method fehlt in {input_file}")
            return pd.DataFrame()

        self.print_missing_columns(
            df=df,
            input_file=input_file,
        )

        for segment_id, group in df.groupby("segment_id"):
            ref_rows = group[
                group["method"] == self.REF_METHOD
            ]

            raw_rows = group[
                group["method"] == self.RAW_METHOD
            ]

            if ref_rows.empty:
                print(
                    f"WARNUNG: REF fehlt fuer Segment "
                    f"{segment_id} in {input_file.name}"
                )
                continue

            if raw_rows.empty:
                print(
                    f"WARNUNG: RAW fehlt fuer Segment "
                    f"{segment_id} in {input_file.name}"
                )
                continue

            ref_row = ref_rows.iloc[0]
            raw_row = raw_rows.iloc[0]

            filter_rows = group[
                ~group["method"].isin(
                    [
                        self.REF_METHOD,
                        self.RAW_METHOD,
                    ]
                )
            ]

            for _, filter_row in filter_rows.iterrows():
                record_value = self.safe_get(
                    filter_row,
                    "record",
                )

                if pd.isna(record_value):
                    record_value = self.safe_get(
                        raw_row,
                        "record",
                    )

                row = {
                    "record": record_value,
                    "segment_id": segment_id,
                    "method": self.safe_get(
                        filter_row,
                        "method",
                    ),
                }

                for metric in self.QUALITY_METRICS:
                    row[metric] = self.safe_get(
                        raw_row,
                        metric,
                    )

                for metric in self.ERROR_METRICS:
                    row[f"improvement_{metric}"] = (
                        self.relative_error_improvement(
                            raw=self.safe_get(raw_row, metric),
                            filtered=self.safe_get(filter_row, metric),
                            ref=self.safe_get(ref_row, metric),
                        )
                    )

                for metric in self.F1_METRICS:
                    row[f"improvement_{metric}"] = (
                        self.f1_improvement(
                            raw=self.safe_get(raw_row, metric),
                            filtered=self.safe_get(filter_row, metric),
                        )
                    )

                for metric in self.CORRELATION_METRICS:
                    row[f"improvement_{metric}"] = (
                        self.correlation_improvement(
                            raw=self.safe_get(raw_row, metric),
                            filtered=self.safe_get(filter_row, metric),
                        )
                    )

                for metric in self.SNR_METRICS:
                    row[f"improvement_{metric}"] = (
                        self.snr_improvement(
                            raw=self.safe_get(raw_row, metric),
                            filtered=self.safe_get(filter_row, metric),
                        )
                    )

                results.append(row)

        result_df = pd.DataFrame(results)

        if result_df.empty:
            return result_df

        expected_columns = [
            "record",
            "segment_id",
            "method",
        ]

        expected_columns += self.QUALITY_METRICS

        expected_columns += [
            f"improvement_{metric}"
            for metric in self.ERROR_METRICS
        ]

        expected_columns += [
            f"improvement_{metric}"
            for metric in self.F1_METRICS
        ]

        expected_columns += [
            f"improvement_{metric}"
            for metric in self.CORRELATION_METRICS
        ]

        expected_columns += [
            f"improvement_{metric}"
            for metric in self.SNR_METRICS
        ]

        for column in expected_columns:
            if column not in result_df.columns:
                result_df[column] = np.nan

        result_df = result_df[expected_columns]
        result_df = self.add_segment_quality_score(result_df)
        result_df = self.add_group_scores(result_df)

        return result_df

    def build_score_summary(
        self,
        score_df: pd.DataFrame,
    ) -> pd.DataFrame:
        if score_df.empty:
            return pd.DataFrame()

        if "segment_quality_score" not in score_df.columns:
            return pd.DataFrame()

        summary_df = score_df.copy()

        try:
            summary_df["quality_bin"] = pd.qcut(
                summary_df["segment_quality_score"],
                q=self.N_SUMMARY_BINS,
                duplicates="drop",
            )
        except ValueError:
            summary_df["quality_bin"] = "all"

        result = summary_df.groupby(
            [
                "method",
                "quality_bin",
            ],
            observed=True,
        ).agg(
            n_segments=("segment_id", "count"),
            median_segment_quality_score=(
                "segment_quality_score",
                "median",
            ),
            mean_medical_score=("medical_score", "mean"),
            mean_morphology_score=("morphology_score", "mean"),
            mean_fidelity_score=("fidelity_score", "mean"),
            mean_overall_score=("overall_score", "mean"),
            median_medical_score=("medical_score", "median"),
            median_morphology_score=("morphology_score", "median"),
            median_fidelity_score=("fidelity_score", "median"),
            median_overall_score=("overall_score", "median"),
            median_worst_improvement=(
                "worst_improvement",
                "median",
            ),
            mean_negative_metric_count=(
                "negative_metric_count",
                "mean",
            ),
            mean_strong_negative_metric_count=(
                "strong_negative_metric_count",
                "mean",
            ),
        ).reset_index()

        result["quality_bin"] = result["quality_bin"].astype(str)

        return result

    def build_decision_analysis(
        self,
        score_df: pd.DataFrame,
    ) -> pd.DataFrame:
        if score_df.empty:
            return pd.DataFrame()

        if "segment_quality_score" not in score_df.columns:
            return pd.DataFrame()

        analysis_df = score_df.copy()

        try:
            analysis_df["decision_bin"] = pd.qcut(
                analysis_df["segment_quality_score"],
                q=self.N_DECISION_BINS,
                duplicates="drop",
            )
        except ValueError:
            analysis_df["decision_bin"] = "all"

        result = analysis_df.groupby(
            [
                "method",
                "decision_bin",
            ],
            observed=True,
        ).agg(
            n_segments=("segment_id", "count"),
            mean_quality_score=("segment_quality_score", "mean"),
            median_quality_score=("segment_quality_score", "median"),
            mean_overall_score=("overall_score", "mean"),
            median_overall_score=("overall_score", "median"),
            mean_worst_improvement=("worst_improvement", "mean"),
            median_worst_improvement=("worst_improvement", "median"),
            mean_negative_metric_count=(
                "negative_metric_count",
                "mean",
            ),
            mean_strong_negative_metric_count=(
                "strong_negative_metric_count",
                "mean",
            ),
        ).reset_index()

        result["decision_bin"] = result["decision_bin"].astype(str)

        return result

    def add_excel_charts(
        self,
        workbook,
    ):
        sheet_name = "decision_analysis"

        if sheet_name not in workbook.sheetnames:
            return

        worksheet = workbook[sheet_name]

        if worksheet.max_row < 2:
            return

        header_map = {
            str(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        required_headers = [
            "method",
            "mean_quality_score",
            "mean_overall_score",
        ]

        if any(
            header not in header_map
            for header in required_headers
        ):
            return

        method_column = header_map["method"]
        quality_column = header_map["mean_quality_score"]
        overall_column = header_map["mean_overall_score"]

        methods = []

        for row_number in range(2, worksheet.max_row + 1):
            method = worksheet.cell(
                row=row_number,
                column=method_column,
            ).value

            if method is not None and method not in methods:
                methods.append(method)

        if not methods:
            return

        chart = ScatterChart()
        chart.title = (
            "Filterwirkung in Abhaengigkeit "
            "von der Signalqualitaet"
        )
        chart.style = 13
        chart.height = 12
        chart.width = 22
        chart.x_axis.title = (
            "Mean Signal Quality Score "
            "(klein = gut, gross = schlecht)"
        )
        chart.y_axis.title = (
            "Mean Overall Improvement "
            "(positiv = Verbesserung)"
        )
        chart.legend.position = "r"
        chart.x_axis.scaling.min = 0.0
        chart.x_axis.scaling.max = 1.0
        chart.y_axis.crosses_at = 0.0

        method_colors = {
            "Bandpass": "1F4E78",
            "SWT": "C55A11",
            "Wavelet": "70AD47",
        }

        for method in methods:
            row_numbers = []

            for row_number in range(2, worksheet.max_row + 1):
                row_method = worksheet.cell(
                    row=row_number,
                    column=method_column,
                ).value

                quality_value = worksheet.cell(
                    row=row_number,
                    column=quality_column,
                ).value

                overall_value = worksheet.cell(
                    row=row_number,
                    column=overall_column,
                ).value

                if row_method != method:
                    continue

                if quality_value is None:
                    continue

                if overall_value is None:
                    continue

                row_numbers.append(row_number)

            if not row_numbers:
                continue

            first_row = min(row_numbers)
            last_row = max(row_numbers)

            x_values = Reference(
                worksheet,
                min_col=quality_column,
                min_row=first_row,
                max_row=last_row,
            )

            y_values = Reference(
                worksheet,
                min_col=overall_column,
                min_row=first_row,
                max_row=last_row,
            )

            series = Series(
                y_values,
                x_values,
                title=str(method),
            )

            color = method_colors.get(
                str(method),
                "5B9BD5",
            )

            series.marker.symbol = "circle"
            series.marker.size = 7
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 24000
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color

            chart.series.append(series)

        if not chart.series:
            return

        chart_anchor_row = worksheet.max_row + 3

        worksheet.add_chart(
            chart,
            f"A{chart_anchor_row}",
        )

    def save_analysis_file(
        self,
        result_df: pd.DataFrame,
        score_df: pd.DataFrame,
        score_summary_df: pd.DataFrame,
        decision_analysis_df: pd.DataFrame,
        sqi_analysis_df: pd.DataFrame,
        input_file: Path,
    ) -> Path:
        analysis_dir = input_file.parent.parent / "analysis"

        analysis_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_file = (
            analysis_dir
            / f"{input_file.stem}_improvements.xlsx"
        )

        with pd.ExcelWriter(
            output_file,
            engine="openpyxl",
        ) as writer:
            result_df.to_excel(
                writer,
                sheet_name="improvements",
                index=False,
            )

            score_df.to_excel(
                writer,
                sheet_name="scores",
                index=False,
            )

            score_summary_df.to_excel(
                writer,
                sheet_name="score_summary",
                index=False,
            )

            decision_analysis_df.to_excel(
                writer,
                sheet_name="decision_analysis",
                index=False,
            )

            sqi_analysis_df.to_excel(
                writer,
                sheet_name="sqi_analysis",
                index=False,
            )

        self.format_excel_file(output_file)

        print(f"Gespeichert: {output_file}")

        return output_file

    def run(self) -> list[Path]:
        files = self.get_input_files()

        if len(files) == 0:
            print(
                "Keine Eingabedateien gefunden unter: "
                f"{self.input_root}/{self.input_pattern}"
            )
            return []

        output_files = []

        print()
        print("Starte separate Record-Analysen")
        print(f"Gefundene Eingabedateien: {len(files)}")
        print()

        for input_file in files:
            print(f"Analysiere: {input_file}")

            df = pd.read_excel(
                input_file,
                engine="openpyxl",
            )

            result_df = self.build_improvement_table(
                df=df,
                input_file=input_file,
            )

            if result_df.empty:
                print(
                    f"Keine Ergebniszeilen fuer: {input_file}"
                )
                continue

            score_df = self.build_score_table(
                result_df=result_df,
                include_source_columns=False,
            )

            score_summary_df = self.build_score_summary(
                score_df
            )

            decision_analysis_df = self.build_decision_analysis(
                score_df
            )

            sqi_analysis_df = self.build_sqi_analysis(
                result_df=result_df,
                group_by_method=True,
            )

            output_file = self.save_analysis_file(
                result_df=result_df,
                score_df=score_df,
                score_summary_df=score_summary_df,
                decision_analysis_df=decision_analysis_df,
                sqi_analysis_df=sqi_analysis_df,
                input_file=input_file,
            )

            output_files.append(output_file)

        print()
        print("Separate Record-Analysen abgeschlossen.")
        print(f"Erzeugte Dateien: {len(output_files)}")

        return output_files
