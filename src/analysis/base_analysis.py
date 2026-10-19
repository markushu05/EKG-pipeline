from pathlib import Path

import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


class BaseAnalysis:
    """
    Basisklasse für die segmentweise und globale Analyse.

    Enthält:
        - gemeinsame Metrikdefinitionen
        - Improvement-Berechnungen
        - Score-Berechnungen
        - SQI-Analyse
        - Excel-Formatierung
    """

    REF_METHOD = "REF"
    RAW_METHOD = "RAW"

    EPSILON = 1e-12

    QUALITY_METRICS = [
        "quality_beat_template_correlation",
        "quality_bsqi",
        "quality_baseline_drift",
        "quality_estimated_snr",
        "quality_rr_plausibility",
    ]

    HIGHER_IS_BETTER_QUALITY_METRICS = [
        "quality_beat_template_correlation",
        "quality_bsqi",
        "quality_estimated_snr",
        "quality_rr_plausibility",
    ]

    LOWER_IS_BETTER_QUALITY_METRICS = [
        "quality_baseline_drift",
    ]

    ERROR_METRICS = [
        "medical_hr_error_bpm",
        "medical_rr_error_ms",
        "morphology_qrs_amplitude_error",
        "morphology_qrs_duration_error",
        "morphology_st_segment_error",
        "morphology_p_amplitude_error",
        "morphology_t_amplitude_error",
        "morphology_qt_interval_error",
        "fidelity_rmse",
        "fidelity_prd",
    ]

    F1_METRICS = [
        "medical_sensitivity",
        "medical_precision",
        "medical_f1_score",
    ]

    CORRELATION_METRICS = [
        "morphology_beat_correlation",
        "fidelity_full_segment_corr",
    ]

    SNR_METRICS = [
        "fidelity_snr_after_db",
    ]

    MEDICAL_IMPROVEMENT_COLUMNS = [
        "improvement_medical_hr_error_bpm",
        "improvement_medical_rr_error_ms",
        "improvement_medical_sensitivity",
        "improvement_medical_precision",
        "improvement_medical_f1_score",
    ]

    MORPHOLOGY_IMPROVEMENT_COLUMNS = [
        "improvement_morphology_beat_correlation",
        "improvement_morphology_qrs_amplitude_error",
        "improvement_morphology_qrs_duration_error",
        "improvement_morphology_st_segment_error",
        "improvement_morphology_p_amplitude_error",
        "improvement_morphology_t_amplitude_error",
        "improvement_morphology_qt_interval_error",
    ]

    FIDELITY_IMPROVEMENT_COLUMNS = [
        "improvement_fidelity_rmse",
        "improvement_fidelity_prd",
        "improvement_fidelity_snr_after_db",
        "improvement_fidelity_full_segment_corr",
    ]

    def __init__(
        self,
        input_root: str | Path = "output",
        input_pattern: str = "*/segmented/results_*.xlsx",
    ):
        self.input_root = Path(input_root)
        self.input_pattern = input_pattern

    @staticmethod
    def safe_get(row: pd.Series, column: str):
        if column not in row.index:
            return np.nan

        return row[column]

    def relative_error_improvement(
        self,
        raw,
        filtered,
        ref,
    ):
        if pd.isna(raw):
            return np.nan

        if pd.isna(filtered):
            return np.nan

        if pd.isna(ref):
            return np.nan

        raw_distance = abs(raw - ref)
        filtered_distance = abs(filtered - ref)

        if raw_distance < self.EPSILON:
            return 0.0

        return (
            raw_distance - filtered_distance
        ) / raw_distance

    def f1_improvement(
        self,
        raw,
        filtered,
    ):
        if pd.isna(raw):
            return np.nan

        if pd.isna(filtered):
            return np.nan

        raw_error = 1.0 - raw
        filtered_error = 1.0 - filtered

        if raw_error < self.EPSILON:
            return 0.0

        return (
            raw_error - filtered_error
        ) / raw_error

    def correlation_improvement(
        self,
        raw,
        filtered,
    ):
        if pd.isna(raw):
            return np.nan

        if pd.isna(filtered):
            return np.nan

        raw_error = 1.0 - raw
        filtered_error = 1.0 - filtered

        if raw_error < self.EPSILON:
            return 0.0

        return (
            raw_error - filtered_error
        ) / raw_error

    def snr_improvement(
        self,
        raw,
        filtered,
    ):
        if pd.isna(raw):
            return np.nan

        if pd.isna(filtered):
            return np.nan

        if abs(raw) < self.EPSILON:
            return 0.0

        return (
            filtered - raw
        ) / abs(raw)

    def get_input_files(self):

        files = sorted(
            self.input_root.glob(
                self.input_pattern
            )
        )

        files = [
            file
            for file in files
            if "_improvements" not in file.name
            and "_analysis" not in file.name
        ]

        return files

    def print_missing_columns(
        self,
        df: pd.DataFrame,
        input_file: Path,
    ):
        required_columns = [
            "record",
            "segment_id",
            "method",
        ]

        required_columns += self.QUALITY_METRICS
        required_columns += self.ERROR_METRICS
        required_columns += self.F1_METRICS
        required_columns += self.CORRELATION_METRICS
        required_columns += self.SNR_METRICS

        missing_columns = [
            column
            for column in required_columns
            if column not in df.columns
        ]

        if len(missing_columns) == 0:
            return

        print()
        print("WARNUNG: Fehlende Spalten in Datei:")
        print(input_file)

        for column in missing_columns:
            print(f"  - {column}")

        print()

    def add_group_scores(
        self,
        result_df: pd.DataFrame,
    ) -> pd.DataFrame:
        result_df = result_df.copy()

        medical_columns = [
            column
            for column in self.MEDICAL_IMPROVEMENT_COLUMNS
            if column in result_df.columns
        ]

        morphology_columns = [
            column
            for column in self.MORPHOLOGY_IMPROVEMENT_COLUMNS
            if column in result_df.columns
        ]

        fidelity_columns = [
            column
            for column in self.FIDELITY_IMPROVEMENT_COLUMNS
            if column in result_df.columns
        ]

        all_improvement_columns = (
            medical_columns
            + morphology_columns
            + fidelity_columns
        )

        if medical_columns:
            result_df["medical_score"] = result_df[
                medical_columns
            ].mean(
                axis=1,
                skipna=True,
            )
        else:
            result_df["medical_score"] = np.nan

        if morphology_columns:
            result_df["morphology_score"] = result_df[
                morphology_columns
            ].mean(
                axis=1,
                skipna=True,
            )
        else:
            result_df["morphology_score"] = np.nan

        if fidelity_columns:
            result_df["fidelity_score"] = result_df[
                fidelity_columns
            ].mean(
                axis=1,
                skipna=True,
            )
        else:
            result_df["fidelity_score"] = np.nan

        result_df["overall_score"] = result_df[
            [
                "medical_score",
                "morphology_score",
                "fidelity_score",
            ]
        ].mean(
            axis=1,
            skipna=True,
        )

        if all_improvement_columns:
            result_df["worst_improvement"] = result_df[
                all_improvement_columns
            ].min(
                axis=1,
                skipna=True,
            )

            result_df["negative_metric_count"] = (
                result_df[all_improvement_columns] < -0.02
            ).sum(axis=1)

            result_df["strong_negative_metric_count"] = (
                result_df[all_improvement_columns] < -0.30
            ).sum(axis=1)

        else:
            result_df["worst_improvement"] = np.nan
            result_df["negative_metric_count"] = 0
            result_df["strong_negative_metric_count"] = 0

        return result_df

    def build_score_table(
        self,
        result_df: pd.DataFrame,
        include_source_columns: bool = False,
    ) -> pd.DataFrame:
        columns = ["record"]

        if include_source_columns:
            columns += [
                "source_file",
                "source_group",
            ]

        columns += [
            "segment_id",
            "method",
            "segment_quality_score",
            "medical_score",
            "morphology_score",
            "fidelity_score",
            "overall_score",
            "worst_improvement",
            "negative_metric_count",
            "strong_negative_metric_count",
        ]

        existing_columns = [
            column
            for column in columns
            if column in result_df.columns
        ]

        return result_df[
            existing_columns
        ].copy()

    def build_sqi_analysis(
        self,
        result_df: pd.DataFrame,
        group_by_method: bool = True,
    ) -> pd.DataFrame:
        rows = []

        if "overall_score" not in result_df.columns:
            return pd.DataFrame()

        if group_by_method and "method" in result_df.columns:
            methods = sorted(
                result_df["method"]
                .dropna()
                .unique()
            )
        else:
            methods = [None]

        for method in methods:
            if method is None:
                method_df = result_df.copy()
            else:
                method_df = result_df[
                    result_df["method"] == method
                ].copy()

            for sqi in self.QUALITY_METRICS:
                if sqi not in method_df.columns:
                    continue

                analysis_df = method_df[
                    [
                        sqi,
                        "overall_score",
                    ]
                ].replace(
                    [np.inf, -np.inf],
                    np.nan,
                ).dropna()

                if len(analysis_df) < 10:
                    continue

                spearman_corr = analysis_df[
                    sqi
                ].corr(
                    analysis_df["overall_score"],
                    method="spearman",
                )

                pearson_corr = analysis_df[
                    sqi
                ].corr(
                    analysis_df["overall_score"],
                    method="pearson",
                )

                rows.append(
                    {
                        "method": method,
                        "sqi": sqi,
                        "n_segments": len(analysis_df),
                        "spearman_corr": spearman_corr,
                        "pearson_corr": pearson_corr,
                        "absolute_spearman": abs(
                            spearman_corr
                        ),
                    }
                )

        result = pd.DataFrame(rows)

        if result.empty:
            return result

        return result.sort_values(
            by=[
                "method",
                "absolute_spearman",
            ],
            ascending=[
                True,
                False,
            ],
            na_position="last",
        )

    @staticmethod
    def make_excel_table_name(
        sheet_name: str,
        index: int,
    ) -> str:
        cleaned_name = "".join(
            character
            if character.isalnum()
            else "_"
            for character in sheet_name
        )

        return f"Table_{index}_{cleaned_name}"[:30]

    def format_excel_file(
        self,
        output_file: Path,
    ):
        workbook = load_workbook(output_file)

        header_fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid",
        )

        header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        for sheet_index, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):
            worksheet.freeze_panes = "A2"

            max_row = worksheet.max_row
            max_column = worksheet.max_column

            if max_row < 1 or max_column < 1:
                continue

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            for column_cells in worksheet.columns:
                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:
                    if cell.value is not None:
                        max_length = max(
                            max_length,
                            len(str(cell.value)),
                        )

                    if isinstance(cell.value, float):
                        cell.number_format = "0.0000"

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    35,
                )

            if max_row < 2:
                continue

            last_column_letter = get_column_letter(
                max_column
            )

            table_range = (
                f"A1:"
                f"{last_column_letter}"
                f"{max_row}"
            )

            table_name = self.make_excel_table_name(
                worksheet.title,
                sheet_index,
            )

            table = Table(
                displayName=table_name,
                ref=table_range,
            )

            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            try:
                worksheet.add_table(table)
            except ValueError:
                pass

        self.add_excel_charts(workbook)

        workbook.save(output_file)

    def add_excel_charts(self, workbook):
        """
        Wird von den Unterklassen überschrieben.
        """
        return