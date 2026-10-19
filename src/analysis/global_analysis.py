from pathlib import Path

import numpy as np
import pandas as pd

from openpyxl.chart import Reference, ScatterChart, Series

from src.analysis.base_analysis import BaseAnalysis


class GlobalBandpassAnalysis(BaseAnalysis):
    """
    Globale Analyse aller Records für den Bandpassfilter.

    Die Quality-Metriken werden gemeinsam über alle Records
    normalisiert.

    Ausgabe:
        output/analysis/all_records_bandpass_analysis.xlsx
    """

    FILTER_METHOD = "Bandpass"

    N_DECISION_BINS = 10

    QUALITY_WEIGHTS = {
        "quality_estimated_snr_score": 0.70,
        "quality_rr_plausibility_score": 0.0,
        "quality_beat_template_correlation_score": 0.30,
        "quality_baseline_drift_score": 0.0,
        "quality_bsqi_score": 0.0,
    }

    def __init__(
        self,
        input_root: str | Path = "output",
        input_pattern: str = "*/segmented/results_*.xlsx",
        output_file: str | Path | None = None,
    ):
        super().__init__(
            input_root=input_root,
            input_pattern=input_pattern,
        )

        if output_file is None:
            self.output_file = (
                self.input_root
                / "analysis"
                / "all_records_bandpass_analysis.xlsx"
            )
        else:
            self.output_file = Path(output_file)

    def build_file_improvement_table(
        self,
        df: pd.DataFrame,
        input_file: Path,
    ) -> pd.DataFrame:
        results = []

        if "segment_id" not in df.columns:
            print(
                f"FEHLER: segment_id fehlt in {input_file}"
            )
            return pd.DataFrame()

        if "method" not in df.columns:
            print(
                f"FEHLER: method fehlt in {input_file}"
            )
            return pd.DataFrame()

        self.print_missing_columns(
            df=df,
            input_file=input_file,
        )

        for segment_id, group in df.groupby(
            "segment_id"
        ):
            ref_rows = group[
                group["method"] == self.REF_METHOD
            ]

            raw_rows = group[
                group["method"] == self.RAW_METHOD
            ]

            filter_rows = group[
                group["method"] == self.FILTER_METHOD
            ]

            if ref_rows.empty:
                continue

            if raw_rows.empty:
                continue

            if filter_rows.empty:
                continue

            ref_row = ref_rows.iloc[0]
            raw_row = raw_rows.iloc[0]
            filter_row = filter_rows.iloc[0]

            record_value = self.safe_get(
                filter_row,
                "record",
            )

            if pd.isna(record_value):
                record_value = self.safe_get(
                    raw_row,
                    "record",
                )

            row = {
                "record": record_value,
                "source_file": input_file.name,
                "source_group": (
                    input_file.parent.parent.name
                ),
                "segment_id": segment_id,
                "method": self.FILTER_METHOD,
            }

            for metric in self.QUALITY_METRICS:
                row[metric] = self.safe_get(
                    raw_row,
                    metric,
                )

            for metric in self.ERROR_METRICS:
                row[
                    f"improvement_{metric}"
                ] = self.relative_error_improvement(
                    raw=self.safe_get(
                        raw_row,
                        metric,
                    ),
                    filtered=self.safe_get(
                        filter_row,
                        metric,
                    ),
                    ref=self.safe_get(
                        ref_row,
                        metric,
                    ),
                )

            for metric in self.F1_METRICS:
                row[
                    f"improvement_{metric}"
                ] = self.f1_improvement(
                    raw=self.safe_get(
                        raw_row,
                        metric,
                    ),
                    filtered=self.safe_get(
                        filter_row,
                        metric,
                    ),
                )

            for metric in self.CORRELATION_METRICS:
                row[
                    f"improvement_{metric}"
                ] = self.correlation_improvement(
                    raw=self.safe_get(
                        raw_row,
                        metric,
                    ),
                    filtered=self.safe_get(
                        filter_row,
                        metric,
                    ),
                )

            for metric in self.SNR_METRICS:
                row[
                    f"improvement_{metric}"
                ] = self.snr_improvement(
                    raw=self.safe_get(
                        raw_row,
                        metric,
                    ),
                    filtered=self.safe_get(
                        filter_row,
                        metric,
                    ),
                )

            results.append(row)

        return pd.DataFrame(results)

    def add_global_segment_quality_score(
        self,
        result_df: pd.DataFrame,
    ) -> pd.DataFrame:
        """
        Normalisiert die RAW-SQIs global über alle Records.

        Ergebnis:
            0 = gute Signalqualität
            1 = schlechte Signalqualität
        """

        result_df = result_df.copy()

        score_columns = []

        for metric in self.QUALITY_METRICS:
            if metric not in result_df.columns:
                continue

            score_column = f"{metric}_score"

            percentile_rank = result_df[
                metric
            ].rank(
                pct=True,
                method="average",
                na_option="keep",
            )

            if metric in self.HIGHER_IS_BETTER_QUALITY_METRICS:
                result_df[score_column] = (
                    1.0 - percentile_rank
                )

            elif metric in self.LOWER_IS_BETTER_QUALITY_METRICS:
                result_df[score_column] = percentile_rank

            else:
                raise ValueError(
                    "Keine Qualitätsrichtung definiert für: "
                    f"{metric}"
                )

            score_columns.append(score_column)

        if len(score_columns) == 0:
            insert_position = min(
                5,
                len(result_df.columns),
            )

            result_df.insert(
                insert_position,
                "segment_quality_score",
                np.nan,
            )

            return result_df

        weighted_sum = pd.Series(
            0.0,
            index=result_df.index,
            dtype=float,
        )

        available_weight = pd.Series(
            0.0,
            index=result_df.index,
            dtype=float,
        )

        for score_column in score_columns:
            weight = self.QUALITY_WEIGHTS.get(
                score_column,
                0.0,
            )

            if weight <= 0.0:
                continue

            valid_mask = result_df[
                score_column
            ].notna()

            weighted_sum.loc[valid_mask] += (
                result_df.loc[
                    valid_mask,
                    score_column,
                ]
                * weight
            )

            available_weight.loc[valid_mask] += weight

        segment_quality_score = pd.Series(
            np.nan,
            index=result_df.index,
            dtype=float,
        )

        usable_mask = (
            available_weight > self.EPSILON
        )

        segment_quality_score.loc[usable_mask] = (
            weighted_sum.loc[usable_mask]
            / available_weight.loc[usable_mask]
        )

        insert_position = min(
            5,
            len(result_df.columns),
        )

        result_df.insert(
            insert_position,
            "segment_quality_score",
            segment_quality_score,
        )

        return result_df

    def build_global_decision_analysis(
        self,
        score_df: pd.DataFrame,
    ) -> pd.DataFrame:
        if score_df.empty:
            return pd.DataFrame()

        analysis_df = score_df.copy()

        try:
            analysis_df["decision_bin"] = pd.qcut(
                analysis_df["segment_quality_score"],
                q=self.N_DECISION_BINS,
                duplicates="drop",
            )

        except ValueError:
            analysis_df["decision_bin"] = "all"

        result = analysis_df.groupby(
            "decision_bin",
            observed=True,
        ).agg(
            n_segments=(
                "segment_id",
                "count",
            ),
            n_records=(
                "record",
                "nunique",
            ),
            mean_quality_score=(
                "segment_quality_score",
                "mean",
            ),
            median_quality_score=(
                "segment_quality_score",
                "median",
            ),
            mean_medical_score=(
                "medical_score",
                "mean",
            ),
            mean_morphology_score=(
                "morphology_score",
                "mean",
            ),
            mean_fidelity_score=(
                "fidelity_score",
                "mean",
            ),
            mean_overall_score=(
                "overall_score",
                "mean",
            ),
            median_overall_score=(
                "overall_score",
                "median",
            ),
            mean_worst_improvement=(
                "worst_improvement",
                "mean",
            ),
            median_worst_improvement=(
                "worst_improvement",
                "median",
            ),
            mean_negative_metric_count=(
                "negative_metric_count",
                "mean",
            ),
            mean_strong_negative_metric_count=(
                "strong_negative_metric_count",
                "mean",
            ),
        ).reset_index()

        result.insert(
            0,
            "method",
            self.FILTER_METHOD,
        )

        result["decision_bin"] = result[
            "decision_bin"
        ].astype(str)

        return result

    def build_global_medical_analysis(
        self,
        result_df: pd.DataFrame,
    ) -> pd.DataFrame:
        if result_df.empty:
            return pd.DataFrame()

        required_columns = [
            "segment_quality_score",
            "segment_id",
            "record",
        ]

        required_columns += (
            self.MEDICAL_IMPROVEMENT_COLUMNS
        )

        missing_columns = [
            column
            for column in required_columns
            if column not in result_df.columns
        ]

        if missing_columns:
            print(
                "Medical Analysis nicht möglich. "
                "Fehlende Spalten:"
            )

            for column in missing_columns:
                print(f"  - {column}")

            return pd.DataFrame()

        analysis_df = result_df.copy()

        try:
            analysis_df["decision_bin"] = pd.qcut(
                analysis_df["segment_quality_score"],
                q=self.N_DECISION_BINS,
                duplicates="drop",
            )

        except ValueError:
            analysis_df["decision_bin"] = "all"

        result = analysis_df.groupby(
            "decision_bin",
            observed=True,
        ).agg(
            n_segments=(
                "segment_id",
                "count",
            ),
            n_records=(
                "record",
                "nunique",
            ),
            mean_quality_score=(
                "segment_quality_score",
                "mean",
            ),
            median_quality_score=(
                "segment_quality_score",
                "median",
            ),
            mean_hr_error_improvement=(
                "improvement_medical_hr_error_bpm",
                "mean",
            ),
            median_hr_error_improvement=(
                "improvement_medical_hr_error_bpm",
                "median",
            ),
            mean_rr_error_improvement=(
                "improvement_medical_rr_error_ms",
                "mean",
            ),
            median_rr_error_improvement=(
                "improvement_medical_rr_error_ms",
                "median",
            ),
            mean_sensitivity_improvement=(
                "improvement_medical_sensitivity",
                "mean",
            ),
            median_sensitivity_improvement=(
                "improvement_medical_sensitivity",
                "median",
            ),
            mean_precision_improvement=(
                "improvement_medical_precision",
                "mean",
            ),
            median_precision_improvement=(
                "improvement_medical_precision",
                "median",
            ),
            mean_f1_improvement=(
                "improvement_medical_f1_score",
                "mean",
            ),
            median_f1_improvement=(
                "improvement_medical_f1_score",
                "median",
            ),
        ).reset_index()

        result.insert(
            0,
            "method",
            self.FILTER_METHOD,
        )

        result["decision_bin"] = result[
            "decision_bin"
        ].astype(str)

        return result

    @staticmethod
    def add_scatter_chart(
        worksheet,
        x_header: str,
        y_header: str,
        title: str,
        y_title: str,
        color: str,
        anchor: str,
    ):
        header_map = {
            str(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        if x_header not in header_map:
            return

        if y_header not in header_map:
            return

        x_column = header_map[x_header]
        y_column = header_map[y_header]

        chart = ScatterChart()

        chart.title = title
        chart.style = 13
        chart.height = 11
        chart.width = 17

        chart.x_axis.title = (
            "Globaler Signal Quality Score "
            "(klein = gut, groß = schlecht)"
        )

        chart.y_axis.title = y_title

        chart.x_axis.scaling.min = 0.0
        chart.x_axis.scaling.max = 1.0
        chart.y_axis.crosses_at = 0.0
        chart.legend = None

        x_values = Reference(
            worksheet,
            min_col=x_column,
            min_row=2,
            max_row=worksheet.max_row,
        )

        y_values = Reference(
            worksheet,
            min_col=y_column,
            min_row=2,
            max_row=worksheet.max_row,
        )

        series = Series(
            y_values,
            x_values,
            title="Bandpass",
        )

        series.marker.symbol = "circle"
        series.marker.size = 8

        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 24000

        series.marker.graphicalProperties.solidFill = color

        series.marker.graphicalProperties.line.solidFill = (
            color
        )

        chart.series.append(series)

        worksheet.add_chart(
            chart,
            anchor,
        )

    def add_decision_charts(
        self,
        workbook,
    ):
        sheet_name = "decision_analysis"

        if sheet_name not in workbook.sheetnames:
            return

        worksheet = workbook[sheet_name]

        if worksheet.max_row < 2:
            return

        first_chart_row = worksheet.max_row + 3

        chart_definitions = [
            {
                "value_header": "mean_medical_score",
                "title": (
                    "Bandpass: Signalqualität "
                    "gegen Medical Score"
                ),
                "y_title": "Mean Medical Improvement",
                "color": "C00000",
                "anchor": f"A{first_chart_row}",
            },
            {
                "value_header": "mean_morphology_score",
                "title": (
                    "Bandpass: Signalqualität "
                    "gegen Morphology Score"
                ),
                "y_title": "Mean Morphology Improvement",
                "color": "70AD47",
                "anchor": f"J{first_chart_row}",
            },
            {
                "value_header": "mean_fidelity_score",
                "title": (
                    "Bandpass: Signalqualität "
                    "gegen Fidelity Score"
                ),
                "y_title": "Mean Fidelity Improvement",
                "color": "ED7D31",
                "anchor": f"A{first_chart_row + 24}",
            },
            {
                "value_header": "mean_overall_score",
                "title": (
                    "Bandpass: Signalqualität "
                    "gegen Overall Score"
                ),
                "y_title": "Mean Overall Improvement",
                "color": "1F4E78",
                "anchor": f"J{first_chart_row + 24}",
            },
        ]

        for definition in chart_definitions:
            self.add_scatter_chart(
                worksheet=worksheet,
                x_header="mean_quality_score",
                y_header=definition["value_header"],
                title=definition["title"],
                y_title=definition["y_title"],
                color=definition["color"],
                anchor=definition["anchor"],
            )

    def add_medical_charts(
        self,
        workbook,
    ):
        sheet_name = "medical_analysis"

        if sheet_name not in workbook.sheetnames:
            return

        worksheet = workbook[sheet_name]

        if worksheet.max_row < 2:
            return

        first_chart_row = worksheet.max_row + 3

        chart_definitions = [
            {
                "value_header": (
                    "mean_hr_error_improvement"
                ),
                "title": (
                    "Bandpass: Signalqualität gegen "
                    "HR Error Improvement"
                ),
                "y_title": "Mean HR Error Improvement",
                "color": "C00000",
                "anchor": f"A{first_chart_row}",
            },
            {
                "value_header": (
                    "mean_rr_error_improvement"
                ),
                "title": (
                    "Bandpass: Signalqualität gegen "
                    "RR Error Improvement"
                ),
                "y_title": "Mean RR Error Improvement",
                "color": "ED7D31",
                "anchor": f"J{first_chart_row}",
            },
            {
                "value_header": (
                    "mean_sensitivity_improvement"
                ),
                "title": (
                    "Bandpass: Signalqualität gegen "
                    "Sensitivity Improvement"
                ),
                "y_title": "Mean Sensitivity Improvement",
                "color": "70AD47",
                "anchor": f"A{first_chart_row + 24}",
            },
            {
                "value_header": (
                    "mean_precision_improvement"
                ),
                "title": (
                    "Bandpass: Signalqualität gegen "
                    "Precision Improvement"
                ),
                "y_title": "Mean Precision Improvement",
                "color": "5B9BD5",
                "anchor": f"J{first_chart_row + 24}",
            },
            {
                "value_header": (
                    "mean_f1_improvement"
                ),
                "title": (
                    "Bandpass: Signalqualität gegen "
                    "F1 Improvement"
                ),
                "y_title": "Mean F1 Improvement",
                "color": "7030A0",
                "anchor": f"A{first_chart_row + 48}",
            },
        ]

        for definition in chart_definitions:
            self.add_scatter_chart(
                worksheet=worksheet,
                x_header="mean_quality_score",
                y_header=definition["value_header"],
                title=definition["title"],
                y_title=definition["y_title"],
                color=definition["color"],
                anchor=definition["anchor"],
            )

    def add_excel_charts(
        self,
        workbook,
    ):
        self.add_decision_charts(workbook)
        self.add_medical_charts(workbook)

    def save_global_analysis_file(
        self,
        result_df: pd.DataFrame,
        score_df: pd.DataFrame,
        decision_analysis_df: pd.DataFrame,
        medical_analysis_df: pd.DataFrame,
        sqi_analysis_df: pd.DataFrame,
    ) -> Path:
        self.output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        with pd.ExcelWriter(
            self.output_file,
            engine="openpyxl",
        ) as writer:
            result_df.to_excel(
                writer,
                sheet_name="improvements",
                index=False,
            )

            score_df.to_excel(
                writer,
                sheet_name="scores",
                index=False,
            )

            decision_analysis_df.to_excel(
                writer,
                sheet_name="decision_analysis",
                index=False,
            )

            medical_analysis_df.to_excel(
                writer,
                sheet_name="medical_analysis",
                index=False,
            )

            sqi_analysis_df.to_excel(
                writer,
                sheet_name="sqi_analysis",
                index=False,
            )

        self.format_excel_file(
            self.output_file
        )

        print(
            f"Gespeichert: {self.output_file}"
        )

        return self.output_file

    def run(self) -> Path | None:
        files = self.get_input_files()

        if len(files) == 0:
            print(
                "Keine Eingabedateien gefunden unter: "
                f"{self.input_root}/{self.input_pattern}"
            )
            return None

        print()
        print("Starte globale Bandpass-Analyse")
        print(f"Gefundene Eingabedateien: {len(files)}")
        print()

        all_frames = []

        for input_file in files:
            print(f"Lese: {input_file}")

            df = pd.read_excel(
                input_file,
                engine="openpyxl",
            )

            file_result_df = (
                self.build_file_improvement_table(
                    df=df,
                    input_file=input_file,
                )
            )

            if not file_result_df.empty:
                all_frames.append(file_result_df)

        if len(all_frames) == 0:
            print(
                "Keine gültigen Bandpass-Ergebniszeilen "
                "erzeugt."
            )
            return None

        result_df = pd.concat(
            all_frames,
            ignore_index=True,
        )

        result_df = (
            self.add_global_segment_quality_score(
                result_df
            )
        )

        result_df = self.add_group_scores(
            result_df
        )

        score_df = self.build_score_table(
            result_df=result_df,
            include_source_columns=True,
        )

        decision_analysis_df = (
            self.build_global_decision_analysis(
                score_df
            )
        )

        medical_analysis_df = (
            self.build_global_medical_analysis(
                result_df
            )
        )

        sqi_analysis_df = self.build_sqi_analysis(
            result_df=result_df,
            group_by_method=True,
        )

        output_file = self.save_global_analysis_file(
            result_df=result_df,
            score_df=score_df,
            decision_analysis_df=decision_analysis_df,
            medical_analysis_df=medical_analysis_df,
            sqi_analysis_df=sqi_analysis_df,
        )

        print()
        print("Globale Bandpass-Analyse abgeschlossen.")

        return output_file