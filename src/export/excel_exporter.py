from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.styles import Font
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.table import TableStyleInfo
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """
    Exportiert SegmentResult-Objekte oder DataFrames
    als formatierte Excel-Tabelle.
    """

    @staticmethod
    def export_results(
            results,
            output_file,
            sheet_name="segment_results"):
        """
        Exportiert eine Liste von SegmentResult-Objekten.

        Erwartet:
        results = [
            SegmentResult(...),
            SegmentResult(...),
            ...
        ]
        """

        rows = []

        for result in results:

            if hasattr(result, "to_dict"):
                rows.append(
                    result.to_dict()
                )

            elif isinstance(result, dict):
                rows.append(result)

            else:
                raise TypeError(
                    "Jedes Result muss entweder eine to_dict()-Methode haben "
                    "oder ein Dictionary sein."
                )

        df = pd.DataFrame(rows)

        ExcelExporter.export_dataframe(
            df=df,
            output_file=output_file,
            sheet_name=sheet_name
        )

    @staticmethod
    def export_dataframe(
            df,
            output_file,
            sheet_name="segment_results"):
        """
        Exportiert einen DataFrame als formatierte Excel-Tabelle.
        """

        output_file = Path(output_file)

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        df.to_excel(
            output_file,
            sheet_name=sheet_name,
            index=False
        )

        ExcelExporter._format_excel_file(
            output_file=output_file,
            sheet_name=sheet_name
        )

    @staticmethod
    def _format_excel_file(
            output_file,
            sheet_name):
        """
        Formatiert die gespeicherte Excel-Datei.
        """

        wb = load_workbook(output_file)
        ws = wb[sheet_name]

        ws.freeze_panes = "A2"

        max_row = ws.max_row
        max_col = ws.max_column

        if max_row < 1 or max_col < 1:
            wb.save(output_file)
            return

        last_col_letter = get_column_letter(max_col)

        table_range = (
            f"A1:{last_col_letter}{max_row}"
        )

        table = Table(
            displayName="SegmentResults",
            ref=table_range
        )

        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style
        ws.add_table(table)

        ExcelExporter._format_header(ws)
        ExcelExporter._adjust_column_widths(ws)
        ExcelExporter._format_numbers(ws)

        wb.save(output_file)

    @staticmethod
    def _format_header(ws):
        """
        Formatiert die Header-Zeile.
        """

        header_fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        for cell in ws[1]:

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    @staticmethod
    def _adjust_column_widths(ws):
        """
        Passt die Spaltenbreite grob automatisch an.
        """

        for column_cells in ws.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            adjusted_width = min(
                max_length + 2,
                25
            )

            ws.column_dimensions[
                column_letter
            ].width = adjusted_width

    @staticmethod
    def _format_numbers(ws):
        """
        Setzt Zahlenformate für Float-Werte.
        """

        max_row = ws.max_row

        for row in ws.iter_rows(
                min_row=2,
                max_row=max_row):

            for cell in row:

                if isinstance(cell.value, float):

                    cell.number_format = "0.0000"